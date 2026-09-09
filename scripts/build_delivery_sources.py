"""Create bounded, source-traceable hazard excerpts from retained raw originals."""

from __future__ import annotations

import argparse
import json
import zipfile
import sys
from hashlib import sha256
from pathlib import Path

import shapefile
from openpyxl import load_workbook
if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.analysis.delivery_sprint import select_intersecting_source_features


LANDSLIDE_LAYERS = ("g_d_rzone", "g_d_yzone", "g_j_yzone", "g_k_rzone", "g_k_yzone")


def _hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _verify_expanded_member(archive: Path, member: str, expanded: Path) -> None:
    with zipfile.ZipFile(archive) as source:
        archived_sha = sha256(source.read(member)).hexdigest()
    if archived_sha != _hash(expanded):
        raise ValueError(f"expanded raw does not match retained archive member: {member}")


def _write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n")


def _aoi(repo: Path, city_id: str):
    filename = {
        "kyoto_kiyomizu": "aoi_kyoto_kiyomizu_gion_v1.geojson",
        "kyoto_arashiyama": "aoi_kyoto_arashiyama_v1.geojson",
        "fujisawa_enoshima": "aoi_fujisawa_enoshima_v1.geojson",
    }[city_id]
    payload = _load(repo / "inputs/staging/PHASE4-MULTICITY-DATA-ACQUISITION/city_aoi" / filename)
    return payload["features"][0]


def _landslide_features(raw_dir: Path, aoi_feature: dict) -> list[dict]:
    result = []
    for layer_id in LANDSLIDE_LAYERS:
        reader = shapefile.Reader(str(raw_dir / f"{layer_id}.shp"), encoding="utf-8", encodingErrors="replace")
        for index, shape_record in enumerate(reader.iterShapeRecords()):
            record = shape_record.record.as_dict()
            result.append({
                "type": "Feature",
                "geometry": shape_record.shape.__geo_interface__,
                "properties": {
                    "source_feature_id": f"{layer_id}:{index}",
                    "source_layer_id": layer_id,
                    "NATURALC": record["NATURALC"],
                    "KINDC": record["KINDC"],
                    "NEWCITYC": record["NEWCITYC"],
                    "_ablepath_display_only": True,
                    "_ablepath_geometry_transform": "NONE_WHOLE_SOURCE_FEATURE_INTERSECTION_SELECTION",
                },
            })
    selected = select_intersecting_source_features(result, aoi_feature, source_crs="EPSG:6668", coverage_crs="EPSG:4326")
    return sorted(selected, key=lambda item: item["properties"]["source_feature_id"])


def _a40_features(path: Path, aoi_feature: dict) -> list[dict]:
    result = []
    payload = _load(path)
    for index, feature in enumerate(payload["features"]):
        props = feature["properties"]
        result.append({
            "type": "Feature",
            "geometry": feature["geometry"],
            "properties": {
                "source_feature_id": f"A40-20_14:{index}",
                "A40_002": props.get("A40_002"),
                "A40_003": props.get("A40_003"),
                "_ablepath_display_only": True,
                "_ablepath_geometry_transform": "NONE_WHOLE_SOURCE_FEATURE_INTERSECTION_SELECTION",
            },
        })
    return select_intersecting_source_features(result, aoi_feature, source_crs="EPSG:6668", coverage_crs="EPSG:4326")


def _facility_source_fields(repo: Path, raw_root: Path) -> tuple[list[dict], dict[str, dict]]:
    """Bind original XLSX attributes to existing facility IDs without geocoding."""
    records = _load(repo / "inputs/staging/KYOTO-OFFICIAL-PARITY-V1/facility_records.json")["records"]
    sources = {
        "kyoto_emergency_shelters_r80818": (
            raw_root / "human-actions/auto-download-kyoto-emergency-shelters-20260901T1912JST/20260825185105_02指定緊急避難場所水害・土砂災害R80818現在.xlsx",
            "e1c03e4e0f830d10430b6a0fb1e1669886c3708cdc61164e5471cea7f9a4d639",
        ),
        "kyoto_designated_shelters_r80818": (
            raw_root / "human-actions/auto-download-kyoto-shelters-20260901T1910JST/20260825185151_01指定避難所R80818現在.xlsx",
            "a30cf1556380359a1fa8fc69e4641fbfb7c3c099595818ac24c01417c9e28246",
        ),
        "kyoto_public_toilets_00307": (
            raw_root / "phase4-multicity-data-acquisition/raw/shared/kyoto_public_toilets_00307/20260306112900_オープンデータ（公衆トイレ、観光トイレ）.xlsx",
            "a46448d7bb6fe814631b1230e69ce6d825d36ccf0ddb35dc339bc49ed60fc764",
        ),
    }
    source_receipts = {}
    for source_id, (path, digest) in sources.items():
        if _hash(path) != digest:
            raise ValueError(f"retained facility raw SHA-256 mismatch: {source_id}")
        source_receipts[source_id] = {"sha256": digest, "filename": path.name, "read_only": True}

    output = []
    for source_id in ("kyoto_emergency_shelters_r80818", "kyoto_designated_shelters_r80818"):
        path, _ = sources[source_id]
        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        by_number = {str(row[0]): row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None}
        for record in (item for item in records if item["source_id"] == source_id):
            source_number = record["facility_record_id"].split(":", 1)[1]
            row = by_number.get(source_number)
            if row is None:
                raise ValueError(f"facility source row missing: {record['facility_record_id']}")
            if source_id == "kyoto_emergency_shelters_r80818":
                attributes = {
                    "official_number": row[0],
                    "administrative_district": row[1],
                    "community_disaster_group": row[2],
                    "flood_evacuation_target_districts": row[4],
                    "landslide_evacuation_target_districts": row[5],
                }
            else:
                attributes = {
                    "official_number": row[0],
                    "administrative_district": row[1],
                    "community_disaster_group": row[2],
                    "listed_maximum_capacity": row[4],
                }
            output.append({
                "facility_record_id": record["facility_record_id"],
                "source_id": source_id,
                "source_sha256": source_receipts[source_id]["sha256"],
                "source_attributes": attributes,
                "current_operation_status": "UNKNOWN",
                "entrance_status": "UNKNOWN",
                "unlock_status": "UNKNOWN",
                "accessibility_status": "UNKNOWN",
                "step_free_status": "UNKNOWN",
                "disaster_availability_status": "UNKNOWN",
            })

    toilet_path, _ = sources["kyoto_public_toilets_00307"]
    toilet_sheet = load_workbook(toilet_path, read_only=True, data_only=True).worksheets[1]
    toilet_rows = list(toilet_sheet.iter_rows(min_row=3, values_only=True))
    for record in (item for item in records if item["source_id"] == "kyoto_public_toilets_00307"):
        # The legacy normalized name may include a concatenated reading.  The
        # two source-provided numeric coordinates are the exact, unique key;
        # this is an identity join only and never a geocoding operation.
        matches = [row for row in toilet_rows if row[6] == record["longitude"] and row[7] == record["latitude"]]
        if len(matches) != 1:
            raise ValueError(f"toilet source identity is not exact: {record['facility_record_id']}")
        row = matches[0]
        output.append({
            "facility_record_id": record["facility_record_id"],
            "source_id": record["source_id"],
            "source_sha256": source_receipts[record["source_id"]]["sha256"],
            "source_attributes": {
                "official_map_number": row[3],
                "administrative_district": row[4],
                "official_page_url": row[8],
                "listed_opening_hours": row[9],
                "listed_baby_support": row[10],
                "listed_ostomate_support": row[11],
                "listed_wheelchair_support": row[12],
                "listed_western_style": row[13],
                "listed_washlet_support": row[14],
                "listed_fixture_count_text": row[15],
            },
            "current_operation_status": "UNKNOWN",
            "entrance_status": "UNKNOWN",
            "unlock_status": "UNKNOWN",
            "accessibility_status": "UNKNOWN",
            "step_free_status": "UNKNOWN",
            "disaster_availability_status": "UNKNOWN",
        })
    return sorted(output, key=lambda item: item["facility_record_id"]), source_receipts


def build(repo: Path, raw_root: Path) -> dict:
    """Read retained originals and write only bounded normalized excerpts."""
    output_root = repo / "inputs/staging/DELIVERY-SPRINT-V1"
    landslide_dir = raw_root / "kyoto_kiyomizu/kyoto_city_hazard_20260830/g_dosha_00all"
    landslide_zip = raw_root / "phase4-multicity-data-acquisition/raw/kyoto_kiyomizu/kyoto_city_landslide_gis_20260830/g_dosha_00all.zip"
    a40_geojson = raw_root / "fujisawa_enoshima/A40-20_14_GML/A40-20_14.GML/A40-20_14.geojson"
    a40_zip = raw_root / "fujisawa_enoshima/A40-20_14_GML.zip"
    expected = {
        landslide_zip: "ac40e0b7116d81516dc10e279f5ca67ba48bbece68dc1164ec8f1d3e0c095828",
        a40_zip: "6b3192e4ed4f8f28d057e4738ecc0d2e7bef232d6adc3022f2d6aec8375f7479",
    }
    for path, digest in expected.items():
        if _hash(path) != digest:
            raise ValueError(f"retained raw SHA-256 mismatch: {path.name}")
    for layer_id in LANDSLIDE_LAYERS:
        for suffix in (".shp", ".shx", ".dbf"):
            _verify_expanded_member(landslide_zip, f"{layer_id}{suffix}", landslide_dir / f"{layer_id}{suffix}")
    _verify_expanded_member(a40_zip, "A40-20_14.GML/A40-20_14.geojson", a40_geojson)
    counts = {}
    for city_id in ("kyoto_kiyomizu", "kyoto_arashiyama"):
        aoi_feature = _aoi(repo, city_id)
        features = _landslide_features(landslide_dir, aoi_feature)
        _write(output_root / city_id / "landslide_aoi_selection.geojson", {
            "type": "FeatureCollection",
            "name": f"{city_id}_official_landslide_aoi_selection",
            "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::6668"}},
            "features": features,
        })
        counts[f"{city_id}_landslide"] = len(features)
    aoi_feature = _aoi(repo, "fujisawa_enoshima")
    features = _a40_features(a40_geojson, aoi_feature)
    _write(output_root / "fujisawa_enoshima/tsunami_a40_aoi_selection.geojson", {
        "type": "FeatureCollection",
        "name": "fujisawa_enoshima_official_tsunami_a40_aoi_selection",
        "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::6668"}},
        "features": features,
    })
    counts["fujisawa_enoshima_tsunami"] = len(features)
    facility_fields, facility_sources = _facility_source_fields(repo, raw_root)
    _write(output_root / "kyoto_facility_source_fields.json", {
        "schema_version": "1.0.0",
        "binding_method": "OFFICIAL_NUMBER_OR_EXACT_SOURCE_COORDINATES",
        "silent_geocoding": False,
        "current_operation_inferred": False,
        "records": facility_fields,
    })
    counts["kyoto_facility_source_fields"] = len(facility_fields)
    receipt = {
        "schema_version": "1.0.0",
        "task_id": "ABLEPATH-KYOTO-FUJISAWA-DELIVERY-SPRINT-V1",
        "raw_mutated": False,
        "geometry_transform": "NONE_WHOLE_SOURCE_FEATURE_INTERSECTION_SELECTION",
        "source_crs": "EPSG:6668",
        "geojson_coordinate_order": "longitude_latitude",
        "sources": {
            "kyoto_landslide": {"source_id": "kyoto_city_landslide_gis_20260830", "revision": "2026-01-22", "sha256": expected[landslide_zip], "license_scope": "PRIVATE_INTERNAL_AND_PRIVATE_GIT_ONLY"},
            "fujisawa_tsunami": {"source_id": "nlni_a40_2020_kanagawa_tsunami", "revision": "2020", "sha256": expected[a40_zip], "license_scope": "PRIVATE_INTERNAL_AND_PRIVATE_GIT_ONLY"},
            **facility_sources,
        },
        "counts": counts,
        "closure_derived": False,
        "damage_derived": False,
        "debris_derived": False,
    }
    _write(output_root / "source_bindings.json", receipt)
    return receipt


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", type=Path, default=Path.cwd())
    parser.add_argument("--raw-root", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(build(args.repo.resolve(), args.raw_root.resolve()), ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
