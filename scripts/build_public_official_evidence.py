"""Build small public-safe DEM and Fujisawa hazard derivatives from verified local raw files."""

from __future__ import annotations

import argparse
import io
import json
import math
import re
from hashlib import sha256
from pathlib import Path
import sys
import zipfile

from openpyxl import load_workbook
import shapefile
from pyproj import CRS, Transformer
from shapely.geometry import mapping, shape
from shapely.ops import transform

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.analysis.delivery_sprint import select_intersecting_source_features
from src.analysis.official_dem import parse_gsi_dem_gml


OUTER_DEM = ("20260902014842719-001.zip", "6b67adaabe15908189316eb34dd574bfe5854b9db509f29902eab7ecfc3a88db")
DEM_PRODUCTS = {
    "kyoto_kiyomizu": {
        "DEM1A": {"52353692": ("FG-GML-523536-DEM1A-20250908.zip", "1e2115ceb62ebc7155006fd750d30b2bd9caa779cd02f7c604fa537d3f46338c", "b5a644e67fc1d5598e977d9164d898599fd96b78bf6faf1c2760337fffdd0090"), "52354602": ("FG-GML-523546-DEM1A-20250908.zip", "968aa88c36f4c1b1b82f49cf0f64815e85d802f15a247fd8c480d8017666883d", "9e84c0c6d96e1f0fa133534eb2c4139c8d365d67a46b93cef172146a4f08f6a6")},
        "DEM5A": {"52353692": ("FG-GML-523536-DEM5A-20250908.zip", "9669221197f1aa2a5c99bf301098248db1ce0e2e9a40fbc76425b2ce6e350721", "79877d14ae15bd4026feb6594d25986f76de47508fc5f4b897b24fde38ff7fb3"), "52354602": ("FG-GML-523546-DEM5A-20250908.zip", "93f70338fb02a372c12addead0c497c80fb6ce489643a1529e00e935c0afdfb8", "9632a761439b73689eca26d917d6d780999e54234736a7064598b15aefd3289c")},
    },
    "kyoto_arashiyama": {
        "DEM1A": {"52354513": ("FG-GML-523545-DEM1A-20250908.zip", "6c427b6b2ec65418ba9a0f23f4a07af9e27a421761292c6e17b3e434e82cd338", "ba72bc09603bd4352421b51bb9754879f98e2eb5076ba631e69183b828b809c5"), "52354514": ("FG-GML-523545-DEM1A-20250908.zip", "6c427b6b2ec65418ba9a0f23f4a07af9e27a421761292c6e17b3e434e82cd338", "0c7ef46afd15fc22f0b4205d834b9e31e5f485adf3155c81a7a2a56abfd4dca7")},
        "DEM5A": {"52354513": ("FG-GML-523545-DEM5A-20250908.zip", "082e675a7c3a495ba3297790b70f3d97f3a005ce50d34372e8b8a708b3853914", "9c4e6f47861703d3929df75881f30d4da8693bd98158ce01c87b16ebb5943dd8"), "52354514": ("FG-GML-523545-DEM5A-20250908.zip", "082e675a7c3a495ba3297790b70f3d97f3a005ce50d34372e8b8a708b3853914", "114da1a4d6b04893db1957aaf0df686227ba3b93b1938adece3ffd7e902e91a4")},
    },
    "fujisawa_enoshima": {
        "DEM1A": {"52397368": ("FG-GML-523973-DEM1A-20250613.zip", "254e94770789b283fc3e6e0c2c480dba782bde7253bf333dc1857cead9cafdc3", "1a9eea555cc33611baf3d07d5ee32b0f44aa1874f2a36c9c611159da5a129bff")},
        "DEM5A": {"52397368": ("FG-GML-523973-DEM5A-20250620.zip", "334bb93ea00b778f334729a7c144f04451c1a3c3a4d3b87f4de01522b6c3eb9c", "e9e41a9ea830376b4511a09710c8010020c22c539ab22b0eefef1f5ee439daf4")},
    },
}
NODE_PATHS = {
    "kyoto_kiyomizu": "cities/kyoto_kiyomizu/graph/real/candidate_nodes.geojson",
    "kyoto_arashiyama": "cities/kyoto_arashiyama/graph/walk_nodes.real.geojson",
    "fujisawa_enoshima": "cities/fujisawa_enoshima/graph/candidate_walk_nodes.real.geojson",
}
EDGE_PATHS = {
    "kyoto_kiyomizu": "cities/kyoto_kiyomizu/graph/real/candidate_edges.geojson",
    "kyoto_arashiyama": "cities/kyoto_arashiyama/graph/walk_edges.real.geojson",
    "fujisawa_enoshima": "cities/fujisawa_enoshima/graph/candidate_walk_edges.real.geojson",
}
HAZARD_SOURCES = {
    "01_.zip": {"sha256": "03989508e8e715496c700f92e30ac8c3feaaec163d19438aa28b3bcb3738b442", "resource_id": "704a2ee0-0040-4a96-b92a-c8e36b559d3d", "resource_response_sha256": "NOT_RETAINED_THIS_RUN", "source_id": "kanagawa_r7_intensity_distribution", "kind": "INTENSITY_SCENARIO", "definition_sheet": "震度分布", "fields": ["KEY_CODE", "ﾒｯｼｭｺｰﾄﾞ", "LON", "LAT", "JMA", "PGA", "SI"], "definition_fields": ["KEY_CODE", "メッシュコード", "LON", "LAT", "JMA", "PGA", "SI"], "member_revision": "R7_MARCH_2025_SCENARIO_SET"},
    "02_.zip": {"sha256": "91f721e0f37114379d7535f5a9c09eac03ffa9c4d3dc1a2088d6d8ee5f47d023", "resource_id": "1f94e194-1764-46db-baf9-e8b017ae458d", "resource_response_sha256": "e80a05d87a45ebb9b5a72e93958fed0881f0548d540ba96b379ba26b0a9c2e76", "source_id": "kanagawa_r7_liquefaction_distribution", "kind": "LIQUEFACTION_SCENARIO", "definition_sheet": "液状化分布", "fields": ["MeshCode", "bic", "I", "PL", "沈下量S", "液状化層厚"], "member_revision": "R7_MARCH_2025_SCENARIO_SET"},
    "13_.zip": {"sha256": "281f43260b2a0d18f8b4afb4fdbfdba79464692ed175d1d0be3c8f88df6dcaa3", "resource_id": "0511f2b8-db28-4eae-a5a9-83ac58d31fbc", "resource_response_sha256": "e2a722ecab192b0e360c39fa6f7d89cebb70106c510933be778e4fe6ff035559", "source_id": "kanagawa_r7_shaking_susceptibility", "kind": "SHAKING_SUSCEPTIBILITY", "definition_sheet": "ゆれやすさマップ", "fields": ["MeshCode", "lon", "lat", "depth", "PGV600", "I600", "更新AVS30", "更新dl600", "更新Isurf", "更新計測震", "更新順位", "ランク"], "definition_fields": ["MeshCode", "lon", "lat", "depth", "PGV600", "I600", "更新AVS30", "更新ｄｌ600", "更新Isurf", "更新計測震", "更新順位", "ランク"], "member_revision": "R6_UPDATE_2025-02-05"},
    "14_.zip": {"sha256": "fb18d09991f63eae642483334c2eb1657b139cc746aa4c4af083260a070146de", "resource_id": "cd7619e9-b6f8-48ad-a251-f8fa9cc84462", "resource_response_sha256": "5e24aa081789fdd1b10766ebab6ad04f83a705df6fbff81e902d0e21230c63fb", "source_id": "kanagawa_r7_liquefaction_hazard", "kind": "LIQUEFACTION_HAZARD", "definition_sheet": "液状化マップ", "fields": ["No", "MeshCode", "建物棟数", "PL", "沈下量S", "危険度ラ", "沈下量ラ"], "member_revision": "R6_UPDATE_2025-02-15_V01"},
}
DEFINITION = ("data_definition.xlsx", "e197e30502da7aa2d6e9f84c20d18c9aeb8daad457f24c2632615acebbe556f1", "5a584282-4a05-4c0a-b05c-bd75f9dd15f3")
KANAGAWA_DATASET = "https://catalog.opendata.pref.kanagawa.jp/dataset/fdc2ffe1fd3cb572d95d0f954f6c72eb"
KANAGAWA_DATASET_RESPONSE_SHA256 = "62208c40000f8bb6a40ecb693fbf0d6c0c50101a8945d4a89a9193d5092008f3"
KANAGAWA_DEFINITION_RESPONSE_SHA256 = "f6675eb85dde9d8b3bfbff8a0a5a059bca68f43da9800dc5f79bded2424754c8"
GSI_SPEC_SHA256 = "b96f524cf407c7dc73347f79d21a8c152edffaf6a049abdc9d50bb764ee82dbb"
GSI_UPDATE_RESPONSE_SHA256 = "b0ee88c7feafc1f4fba10166322c402e2cae7ca56566fecfe019e41b951d4ebd"
GSI_DATUM_RESPONSE_SHA256 = "a01eb54761186b6fbefd1ee2ad69f10324121deb26c44981d1eb4088ee847be2"
GSI_TERMS_RESPONSE_SHA256 = "5388d2854a05f3e7b41cb980ecfb2de75b315de68e0d9c94b9334c11e580739e"
_COMMON_HAZARD_SHP_SHA256 = "429ab4289988e596003c7d92fe4279c4ca13bea3079190fa8af2e4ac265f3108"
_COMMON_HAZARD_TXT_SHA256 = "1a57d4f1596e49237fd31a40776677576c90acc9cadd0069a574c90d5f4da616"
# The eight intensity scenarios differ only in their DBF payload; SHP and the
# (erroneous) CRS sidecar are byte-identical across members.
_COMMON_INTENSITY_SHP_SHA256 = "1fcd2babd92e6cfd7717688a173672b7ed056be157cc634f6d0891710829fb4a"
_COMMON_INTENSITY_TXT_SHA256 = "fbd6b516c66445b357055a4bf77c0629a16d53b1df08c48650001c0558b02258"
# Provider sibling layer used only to read the provider's own PROJCS declaration
# and to confirm that the sibling carries the identical cell geometry.
INTENSITY_SIBLING = {
    "filename": "00_.zip",
    "sha256": "93b0adea6b7ae45fa3732e4051dbcaeb5fa63bf6f5bd1aa61e041a2679c3c1da",
    "member_suffix": "250mesh_人口_region",
    "declared_projcs": "JGD_2011_Japan_Zone_9",
}
INTENSITY_ERRONEOUS_SIDECAR_DECLARED = "GCS_Tokyo"
# Verification-only comparison tolerance for the EPSG:6677 cross-check. It is not
# registered in data/constants_registry.yaml and is never used by any evaluation,
# score, threshold, or state decision.
INTENSITY_CROSS_CHECK_TOLERANCE_M = 0.05
# T-B METHOD#3, verification only: GRS80 ellipsoidal distance, not a degree box.
INTENSITY_CENTRE_TOLERANCE_M = 0.06
_INTENSITY_GEOD = CRS.from_epsg(6668).get_geod()
# Canonical JSON section bytes independently matched at PR14 f117f8d and f054ed4.
FROZEN_DEM_SECTION_SHA256 = "f25b6d51e1ddbe11e8185dd73ebb016e14c9101592d6aecd0e3e0e806ef71a05"
INTENSITY_LICENSE_NOTE = "CC-BY-4.0; retrospective resource metadata binding: inputs/staging/FUJISAWA-INTENSITY-CRS-V1/license_binding.json; not a backdated acquisition receipt"
HAZARD_MEMBER_CONTRACTS = {
    "01_.zip": [
        ("IntS-01", "三浦半島断層群の地震", "e46387fe06fc0d6cf7fe69eaa83d0700e7e5a729d3f3bda41812214a3da0251d", "b8063014b1da3adc1b9d6232f2467c9f5b1d56b9e3806127aa2ec240a3d380cb"),
        ("IntS-02", "元禄型関東地震", "94e369df55e54e976116f2865b00e87e013e0aa8103d6c3f289b02f69c75ff05", "3c65c7dfe12d394ffef456061086c12c7ad08fa94899e242533956616c7c85d0"),
        ("IntS-03", "南海トラフ巨大地震", "a668c1cc1608d7e4a066e01dadef36eadd58740f1ef99ed06c067b76b2378742", "35b5bfda624f7f04fb4654332e25eb6292f2c38513b2e4af56e439d4ec2632bb"),
        ("IntS-04", "大正型関東地震", "88a8dc3c8c352f20b9966d6b97f36d9103f1c2d85b41ec63d44db87cfeb82d89", "8c9d7cb8f635261e19a9b0fda633ff29de82fbcdc2e24f13600d98b3771fec4b"),
        ("IntS-05", "東海地震", "1814db122e1ad54907f153bf9a304948506340be009c17ceb48b4fb977d60884", "6952d2fe35c22a99fac73d47fa2b5a36efb09f6f1e717ee74b0e12f4fed260d2"),
        ("IntS-06", "相模トラフ沿いの最大クラスの地震", "3d5e81fe355bbdb475d2ebf19f270f1a89a8b06f92d6dfd7d521cc632488ae98", "fb1d7765095d1a540bf02f49adb2185fe1e01a203d4afe8d139035892f865b50"),
        ("IntS-07", "神奈川県西部地震", "208106feeca113e9a28c315f86bfec174f6e95c05969f968834061bd2386e8b2", "0870dc6c5eea1904a748a2606ed244d5e82a9e817c8dfc11f87380e319649a9f"),
        ("IntS-08", "都心南部直下地震", "eaa7f05e423e016e875491004d4e8c647db86529bea6782f0749696e8cd15a5c", "86a9537e19f6591ff2a11b467deb5b73f18b57d1848283c66aeb59d5d6880ca0"),
    ],
    "02_.zip": [
        ("LiqS-01", "都心南部直下地震", "8eea69c0dbf76e4327766f41a618dea96c5c5e50c307faf825e144569789e07f", "c7796652a886ef367a60712abbb8574c44101a893ca5e9e59912d7346aa3d1ec"),
        ("LiqS-02", "三浦半島断層群の地震", "85369d8573ac775b8044c252507bd924a181ab95255be3c55de54d9208c1d244", "27d447c884755b6470552d99655a8b40cd301f6ea7e0b46546b3cd0da816aef8"),
        ("LiqS-03", "神奈川県西部地震", "22210a8a75779742d88a549f7f93200128430f31a7539acaa2f57357078e2926", "6f18f446b112f218b7c5c425bb7377a957228d5139dacae05db31a5108953588"),
        ("LiqS-04", "東海地震", "fdb1ee8592f4e549fa8eafca493340ae5045b3c2b080225e32d7b83913807587", "639dfe36bf74edb3787456c503884c72dafbef01b4ce2e1b50f43fd9ce57925d"),
        ("LiqS-05", "大正型関東地震", "9759490a1cac171a88a91562eb8d71d34266d6cb93c7e3c53ba4d8fccb5c105f", "238fd29ae855c0ef4e6c3de92ddc005a3eacfc30bf7e06156ae7d2c7b24a5c9d"),
        ("LiqS-06", "元禄型関東地震", "920c28eb2fd50c760b1fb7b4ca6b12b17ffc6213167f33c23168e85da2327dd9", "037df1417b7615f368df4e7cf3e059bd01849d30bed9c680af4a556a7f674740"),
        ("LiqS-07", "相模トラフ沿いの最大クラスの地震", "39353350477f7c020f593c52f4d29ab1d12f850b83287e34acbec24800ba679d", "5aefa898d34430de5e1f59c9d5d4f5fc0fae96ad209a648a899ad9a5ad1030ee"),
        ("LiqS-08", "南海トラフ巨大地震", "7d44de4b69932b584252e5c4b1dd291ff0dfb369cebcef5d67cb1da383b449f2", "016fc75c7fc9b7a91b898e95d327fe3ddd9b333487d0ac5e0e14878d482de199"),
    ],
    "13_.zip": [("SHAKING_R6", "令和6年度更新版 ゆれやすさ", "c75453386c378cd7146ac6b6de146297ace58b7dae7376a22ca9fd22e61e7ddc", "68897f5ea3ce88335195fecd63bc46969abf9469912f99fe4ca1559edf94ee60")],
    "14_.zip": [("LIQUEFACTION_HAZARD_R6", "令和6年度更新版 液状化危険度", "7da6f012bbb4d41c0a52415d49fdfce71ee9ab3d0678f16cf99b694515bab195", "6eba93771b177f4cf662c4f3f451552ca140ee6d9b278e7a6a8a53f74b9f0580")],
}


def _digest(payload: bytes) -> str:
    return sha256(payload).hexdigest()


def _verified_bytes(path: Path, expected: str) -> bytes:
    payload = path.read_bytes()
    if _digest(payload) != expected:
        raise ValueError(f"raw SHA-256 mismatch: {path.name}")
    return payload


def _normalized_member_name(archive: zipfile.ZipFile, name: str) -> str:
    """Return the provider's own CP932 member path.

    These archives are written on a CP932 system without the ZIP UTF-8 flag, so
    ``zipfile`` falls back to CP437. Normalizing back makes every member path
    hash identical regardless of the machine that runs the build.
    """
    if archive.getinfo(name).flag_bits & 0x800:
        return name
    try:
        return name.encode("cp437").decode("cp932")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name


def _mesh_from_name(name: str) -> str | None:
    match = re.search(r"FG-GML-(\d{4})-(\d{2})-(\d{2})-", name)
    return "".join(match.groups()) if match else None


def _sample_coordinate(
    *,
    city_id: str,
    dem_class: str,
    longitude: float,
    latitude: float,
    sample_id: str,
    identity_fields: dict,
    members: list[tuple],
) -> dict:
    candidates = []
    for containing_name, containing_sha, member_name, member_sha, grid in members:
        sample = grid.sample_from_epsg4326(longitude, latitude)
        if sample["status"] != "OUTSIDE_MEMBER_COVERAGE":
            candidates.append((containing_name, containing_sha, member_name, member_sha, grid, sample))
    if len(candidates) != 1:
        sample = {"status": "AMBIGUOUS_OR_MISSING_MEMBER", "elevation_m": None, "raw_value": None, "surface_type": None, "reason": f"Expected one member at query coordinate; found {len(candidates)}."}
        source_identity = {"member_name": None, "member_sha256": None, "nested_zip_name": None, "nested_zip_sha256": None, "mesh_id": None, "srs_name": None}
    else:
        containing_name, containing_sha, member_name, member_sha, grid, sample = candidates[0]
        source_identity = {"member_name": member_name, "member_sha256": member_sha, "nested_zip_name": containing_name, "nested_zip_sha256": containing_sha, "mesh_id": grid.mesh_id, "srs_name": grid.srs_name}
    return {
        "sample_id": sample_id,
        "city_id": city_id,
        **identity_fields,
        "product": dem_class,
        "query_longitude": longitude,
        "query_latitude": latitude,
        "query_crs": "EPSG:4326",
        "horizontal_transform": "EPSG:4326_TO_EPSG:6668_NO_BALLPARK; JGD2024 horizontal is the renamed unchanged JGD2011 CRS",
        "vertical_reference": "JGD2024_VERTICAL_JAPAN_DATUM_2024",
        "unit": "m",
        **source_identity,
        **sample,
        "step_inferred": False,
        "cross_slope_inferred": False,
    }


def _build_dem(repo: Path, raw_root: Path) -> dict:
    outer_name, outer_sha = OUTER_DEM
    outer_bytes = _verified_bytes(raw_root / outer_name, outer_sha)
    outer = zipfile.ZipFile(io.BytesIO(outer_bytes))
    cities = {}
    for city_id, products in DEM_PRODUCTS.items():
        nodes = json.loads((repo / NODE_PATHS[city_id]).read_text(encoding="utf-8"))["features"]
        edges = json.loads((repo / EDGE_PATHS[city_id]).read_text(encoding="utf-8"))["features"]
        city_products = []
        city_samples = []
        for dem_class, member_sources in products.items():
            # The outer package contains one product ZIP per 6-digit mesh. A city may
            # need the adjacent product ZIP as well, so locate each required 8-digit
            # member across same-class nested ZIPs and verify the containing ZIP.
            members = []
            for wanted, (nested_name, nested_sha, member_sha) in sorted(member_sources.items()):
                nested_bytes = outer.read(nested_name)
                if _digest(nested_bytes) != nested_sha:
                    raise ValueError(f"nested DEM SHA-256 mismatch: {nested_name}")
                nested = zipfile.ZipFile(io.BytesIO(nested_bytes))
                member_name = next((name for name in nested.namelist() if name.endswith(".xml") and _mesh_from_name(name) == wanted), None)
                if member_name is None:
                    raise ValueError(f"missing DEM member {wanted}/{dem_class}")
                member_bytes = nested.read(member_name)
                if _digest(member_bytes) != member_sha:
                    raise ValueError(f"DEM member SHA-256 mismatch: {member_name}")
                grid = parse_gsi_dem_gml(member_bytes, dem_class=dem_class, revision=member_name.rsplit("-", 1)[-1][:-4])
                members.append((nested_name, nested_sha, member_name, member_sha, grid))
            product_samples = []
            for node in sorted(nodes, key=lambda item: item["properties"]["node_id"]):
                node_id = node["properties"]["node_id"]
                longitude, latitude = node["geometry"]["coordinates"]
                row = _sample_coordinate(city_id=city_id, dem_class=dem_class, longitude=longitude, latitude=latitude, sample_id=f"{city_id}:node:{node_id}:{dem_class}", identity_fields={"sample_role": "GRAPH_NODE", "node_id": node_id}, members=members)
                product_samples.append(row)
                city_samples.append(row)
            for edge in sorted(edges, key=lambda item: item["properties"]["edge_id"]):
                edge_id = edge["properties"]["edge_id"]
                for vertex_index, (longitude, latitude) in enumerate(edge["geometry"]["coordinates"]):
                    row = _sample_coordinate(city_id=city_id, dem_class=dem_class, longitude=longitude, latitude=latitude, sample_id=f"{city_id}:edge:{edge_id}:{vertex_index}:{dem_class}", identity_fields={"sample_role": "EDGE_VERTEX", "edge_id": edge_id, "vertex_index": vertex_index}, members=members)
                    product_samples.append(row)
                    city_samples.append(row)
            city_products.append({
                "dataset_id": f"gsi_{city_id}_{dem_class.lower()}_native_cells_2025",
                "product": dem_class, "dem_class": dem_class,
                "mesh_ids": sorted(member_sources), "horizontal_crs": "JGD2024/(B,L); implemented through unchanged JGD2011 horizontal EPSG:6668",
                "vertical_datum": "JGD2024_VERTICAL_JAPAN_DATUM_2024", "unit": "m",
                "aoi": city_id, "aoi_status": "FULL_MEMBER_LEVEL_NODE_COVERAGE",
                "validation_result": "OFFICIAL_SPEC_AND_RAW_SHA_BOUND_NATIVE_CELL_SAMPLING",
                "license_status": "GSI_TERMS_WITH_ATTRIBUTION", "terrain_connected": True,
                "status": "NATIVE_CELL_SAMPLES_CONNECTED", "connected": True,
                "sample_count": len(product_samples), "sample_record_count": len(product_samples),
                "unique_coordinate_count": len({(row["query_longitude"], row["query_latitude"]) for row in product_samples}),
                "node_sample_count": len(nodes), "edge_vertex_sample_count": sum(len(edge["geometry"]["coordinates"]) for edge in edges),
                "numeric_sample_count": sum(row["elevation_m"] is not None for row in product_samples),
                "null_sample_count": sum(row["elevation_m"] is None for row in product_samples),
                "null_record_count": sum(row["elevation_m"] is None for row in product_samples),
                "null_coordinate_count": len({(row["query_longitude"], row["query_latitude"]) for row in product_samples if row["elevation_m"] is None}),
                "member_count": len(members),
                "members": [
                    {"nested_zip_name": nested_name, "nested_zip_sha256": nested_sha, "member_name": member_name, "member_sha256": member_sha, "mesh_id": grid.mesh_id, "srs_name": grid.srs_name}
                    for nested_name, nested_sha, member_name, member_sha, grid in members
                ],
                "implicit_precedence": False, "mosaic_applied": False,
            })
        cities[city_id] = {
            "status": "NATIVE_CELL_SAMPLES_CONNECTED", "connected": True,
            "sample_record_count": len(city_samples),
            "unique_coordinate_count": len({(row["query_longitude"], row["query_latitude"]) for row in city_samples}),
            "numeric_record_count": sum(row["elevation_m"] is not None for row in city_samples),
            "null_record_count": sum(row["elevation_m"] is None for row in city_samples),
            "null_coordinate_count": len({(row["query_longitude"], row["query_latitude"]) for row in city_samples if row["elevation_m"] is None}),
            "products": city_products, "samples": city_samples,
        }
    return {
        "schema_version": "1.0.0", "source_id": "gsi_fundamental_geospatial_dem",
        "source_url": "https://service.gsi.go.jp/kiban/app/spec_update_info/",
        "specification_url": "https://service.gsi.go.jp/kiban/contents/screen/basismap/documents/FGD_DLFileSpecV5.3.pdf",
        "specification_sha256": GSI_SPEC_SHA256,
        "specification_version": "5.3",
        "specification_sections": ["5.2.3.7 DEM tupleList/startPoint/sequenceRule", "GSI elevation datum revision notice effective 2025-07-31"],
        "specification_observed_at": "2026-09-10",
        "spec_update_response_sha256": GSI_UPDATE_RESPONSE_SHA256,
        "datum_revision_response_sha256": GSI_DATUM_RESPONSE_SHA256,
        "outer_sha256": outer_sha, "raw_root_id": "EXTERNAL_OFFICIAL_DATA_ROOT",
        "horizontal_contract": "GSI states JGD2024 horizontal definitions are unchanged from JGD2011; EPSG:6668 is used with no-ballpark transform from EPSG:4326.",
        "vertical_contract": "The verified package was downloaded after the 2025-07-31 GSI provision change; the official update notice binds provided elevation to JGD2024 even where member source-update dates are earlier. Products remain separate.",
        "license_status": "GSI_TERMS_WITH_ATTRIBUTION",
        "license_url": "https://www.gsi.go.jp/kikakuchousei/kikakuchousei40182.html",
        "license_response_sha256": GSI_TERMS_RESPONSE_SHA256,
        "license_scope": "SMALL_NATIVE_CELL_DERIVED_SAMPLES_WITH_ATTRIBUTION_NOT_RAW_PACKAGE_BLANKET",
        "legal_compliance_claim": False,
        "attribution": "基盤地図情報（国土地理院）を加工して作成",
        "no_interpolation": True, "no_smoothing": True, "no_step_inference": True, "no_cross_slope_inference": True,
        "cities": cities,
    }


def _jis_quarter_mesh_cell(mesh_code: str) -> tuple[float, float, float, float]:
    """Return (lon, lat, dlon, dlat) of a JIS X 0410 10-digit quarter mesh cell in EPSG:6668."""
    if len(mesh_code) != 10 or not mesh_code.isdigit():
        raise ValueError(f"unexpected JIS X 0410 quarter-mesh code: {mesh_code!r}")
    primary, secondary = int(mesh_code[0:2]), int(mesh_code[2:4])
    latitude = primary / 1.5 + int(mesh_code[4]) * 5 / 60 + int(mesh_code[6]) * 30 / 3600
    longitude = 100 + secondary + int(mesh_code[5]) * 7.5 / 60 + int(mesh_code[7]) * 45 / 3600
    delta_latitude, delta_longitude = 30 / 3600, 45 / 3600
    for digit in mesh_code[8:]:
        quadrant = int(digit)
        if quadrant not in (1, 2, 3, 4):
            raise ValueError(f"unexpected JIS X 0410 halving digit: {mesh_code!r}")
        delta_latitude /= 2
        delta_longitude /= 2
        if quadrant in (3, 4):
            latitude += delta_latitude
        if quadrant in (2, 4):
            longitude += delta_longitude
    return longitude, latitude, delta_longitude, delta_latitude


def _intensity_centre_error_m(mesh_code: str, longitude: float, latitude: float) -> float:
    west, south, dx, dy = _jis_quarter_mesh_cell(mesh_code)
    if any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) for value in (longitude, latitude)):
        raise ValueError(f"invalid JIS X 0410 cell centre: {mesh_code}")
    distance = _INTENSITY_GEOD.inv(west + dx / 2, south + dy / 2, longitude, latitude)[2]
    if not math.isfinite(distance) or distance > INTENSITY_CENTRE_TOLERANCE_M:
        raise ValueError(f"intensity record exceeds the JIS X 0410 cell centre 0.06 m contract: {mesh_code}")
    return distance


def _intensity_cell(record: dict) -> tuple[list[tuple[float, float]], str]:
    """Reconstruct the cell ring from KEY_CODE and fail closed if LON/LAT disagree."""
    mesh_code = f"{int(record['KEY_CODE']):d}"
    if mesh_code != f"{int(record['ﾒｯｼｭｺｰﾄﾞ']):d}":
        raise ValueError(f"intensity mesh-code fields disagree: {mesh_code}")
    longitude, latitude, delta_longitude, delta_latitude = _jis_quarter_mesh_cell(mesh_code)
    _intensity_centre_error_m(mesh_code, record["LON"], record["LAT"])
    ring = [
        (longitude, latitude),
        (longitude + delta_longitude, latitude),
        (longitude + delta_longitude, latitude + delta_latitude),
        (longitude, latitude + delta_latitude),
        (longitude, latitude),
    ]
    return ring, mesh_code


def _max_vertex_error_m(corners: list[tuple[float, float]], observed: list[tuple[float, float]]) -> float:
    """Largest distance from an observed projected vertex to its nearest reconstructed corner."""
    return max(
        min(((x - cx) ** 2 + (y - cy) ** 2) ** 0.5 for cx, cy in corners)
        for x, y in observed
    )


def _assert_intensity_cross_check(max_vertex_error_m: float, max_sibling_error_m: float, member_id: str) -> None:
    """Fail closed when the reconstruction disagrees with the provider geometry."""
    if max(max_vertex_error_m, max_sibling_error_m) > INTENSITY_CROSS_CHECK_TOLERANCE_M:
        raise ValueError(f"intensity CRS cross-check exceeded {INTENSITY_CROSS_CHECK_TOLERANCE_M} m for {member_id}")


def _safe_properties(record: dict, kind: str, expected_fields: list[str]) -> dict:
    if list(record) != expected_fields:
        raise ValueError(f"unexpected DBF schema for {kind}: {list(record)!r}")
    if kind == "INTENSITY_SCENARIO":
        return {"mesh_code": f"{int(record['KEY_CODE']):d}", "longitude_raw": record["LON"], "latitude_raw": record["LAT"], "jma_raw": record["JMA"], "pga_raw": record["PGA"], "si_raw": record["SI"]}
    if kind == "LIQUEFACTION_SCENARIO":
        return {"mesh_code": str(record["MeshCode"]), "bic": record["bic"], "intensity_raw": record["I"], "pl_raw": record["PL"], "settlement_m_raw": record["沈下量S"], "liquefaction_layer_thickness_m_raw": record["液状化層厚"]}
    if kind == "SHAKING_SUSCEPTIBILITY":
        output_keys = ("mesh_code", "longitude_raw", "latitude_raw", "depth_raw", "pgv600_raw", "i600_raw", "updated_avs30_raw", "updated_dl600_raw", "updated_isurf_raw", "updated_jma_raw", "updated_rank_number_raw", "rank_raw")
        return dict(zip(output_keys, (record[field] for field in expected_fields)))
    output_keys = ("record_number", "mesh_code", "ground_model_raw", "pl_raw", "settlement_m_raw", "hazard_label_raw", "hazard_rank_raw")
    return dict(zip(output_keys, (record[field] for field in expected_fields)))


def _shape_features(archive: zipfile.ZipFile, shp_name: str, path_sha256: str, kind: str, expected_fields: list[str]) -> tuple[list[dict], dict]:
    base = shp_name[:-4]
    required = {extension: base + extension for extension in (".shp", ".shx", ".dbf", ".txt")}
    missing = [name for name in required.values() if name not in archive.namelist()]
    if missing:
        raise ValueError(f"missing shape member companions: {missing}")
    txt_name = required[".txt"]
    crs_wkt = archive.read(txt_name).decode("cp932")
    source_crs = CRS.from_wkt(crs_wkt)
    if source_crs.to_epsg() != 4612 or source_crs.axis_info[0].unit_name.lower() != "degree":
        raise ValueError(f"unexpected source CRS sidecar: {txt_name}")
    reader = shapefile.Reader(shp=io.BytesIO(archive.read(base + ".shp")), shx=io.BytesIO(archive.read(base + ".shx")), dbf=io.BytesIO(archive.read(base + ".dbf")), encoding="cp932")
    if [field[0] for field in reader.fields[1:]] != expected_fields:
        raise ValueError(f"unexpected DBF fields: {shp_name}")
    result = []
    for index, shape_record in enumerate(reader.iterShapeRecords()):
        props = _safe_properties(shape_record.record.as_dict(), kind, expected_fields)
        props["source_feature_id"] = f"{props['mesh_code']}:{index}"
        props["source_class"] = json.dumps(props, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        result.append({"type": "Feature", "properties": props, "geometry": mapping(shape(shape_record.shape.__geo_interface__))})
    member_receipt = {
        "member_path_sha256": path_sha256,
        "shp_sha256": _digest(archive.read(required[".shp"])),
        "shx_sha256": _digest(archive.read(required[".shx"])),
        "dbf_sha256": _digest(archive.read(required[".dbf"])),
        "txt_sha256": _digest(archive.read(required[".txt"])),
    }
    return result, member_receipt


def _intensity_shape_features(archive: zipfile.ZipFile, shp_name: str, path_sha256: str, expected_fields: list[str]) -> tuple[list[dict], dict, dict[str, list[tuple[float, float]]]]:
    """Read intensity records, rebuilding geometry from KEY_CODE instead of the projected SHP.

    The bundled .txt sidecar declares GCS_Tokyo, which contradicts the projected
    coordinate values it accompanies; it is recorded and never used. The projected
    ring is returned only so the caller can cross-check the reconstruction.
    """
    base = shp_name[:-4]
    required = {extension: base + extension for extension in (".shp", ".shx", ".dbf", ".txt")}
    missing = [name for name in required.values() if name not in archive.namelist()]
    if missing:
        raise ValueError(f"missing shape member companions: {missing}")
    sidecar_text = archive.read(required[".txt"]).decode("cp932")
    if INTENSITY_ERRONEOUS_SIDECAR_DECLARED not in sidecar_text:
        raise ValueError(f"unexpected intensity CRS sidecar: {required['.txt']}")
    reader = shapefile.Reader(shp=io.BytesIO(archive.read(required[".shp"])), shx=io.BytesIO(archive.read(required[".shx"])), dbf=io.BytesIO(archive.read(required[".dbf"])), encoding="cp932")
    if [field[0] for field in reader.fields[1:]] != expected_fields:
        raise ValueError(f"unexpected DBF fields: {shp_name}")
    features = []
    projected_rings = {}
    for index, shape_record in enumerate(reader.iterShapeRecords()):
        record = shape_record.record.as_dict()
        ring, mesh_code = _intensity_cell(record)
        props = _safe_properties(record, "INTENSITY_SCENARIO", expected_fields)
        props["source_feature_id"] = f"{mesh_code}:{index}"
        props["source_class"] = json.dumps(props, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        features.append({"type": "Feature", "properties": props, "geometry": {"type": "Polygon", "coordinates": [[list(vertex) for vertex in ring]]}})
        projected_rings[mesh_code] = [tuple(point) for point in shape_record.shape.points[:-1]]
    member_receipt = {
        "member_path_sha256": path_sha256,
        "shp_sha256": _digest(archive.read(required[".shp"])),
        "shx_sha256": _digest(archive.read(required[".shx"])),
        "dbf_sha256": _digest(archive.read(required[".dbf"])),
        "txt_sha256": _digest(archive.read(required[".txt"])),
    }
    return features, member_receipt, projected_rings


def _intensity_sibling_rings(raw_root: Path) -> tuple[dict[str, list[tuple[float, float]]], str]:
    """Read the provider sibling layer that carries the provider's own PROJCS declaration."""
    raw_bytes = _verified_bytes(raw_root / INTENSITY_SIBLING["filename"], INTENSITY_SIBLING["sha256"])
    archive = zipfile.ZipFile(io.BytesIO(raw_bytes))
    base = next((name[:-4] for name in archive.namelist() if name.endswith(".shp") and _normalized_member_name(archive, name).endswith(INTENSITY_SIBLING["member_suffix"] + ".shp")), None)
    if base is None:
        raise ValueError("missing provider sibling geometry member")
    sidecar_text = archive.read(base + ".txt").decode("cp932")
    if INTENSITY_SIBLING["declared_projcs"] not in sidecar_text:
        raise ValueError("provider sibling sidecar does not declare the expected PROJCS")
    reader = shapefile.Reader(shp=io.BytesIO(archive.read(base + ".shp")), shx=io.BytesIO(archive.read(base + ".shx")), dbf=io.BytesIO(archive.read(base + ".dbf")), encoding="cp932")
    rings = {}
    for shape_record in reader.iterShapeRecords():
        rings[f"{int(shape_record.record.as_dict()['KEY_CODE']):d}"] = [tuple(point) for point in shape_record.shape.points[:-1]]
    return rings, _digest(archive.read(base + ".txt"))


def _bind_hazard_member_inventory(observed: list[dict], contracts: list[tuple], *, shp_sha256: str = _COMMON_HAZARD_SHP_SHA256, txt_sha256: str = _COMMON_HAZARD_TXT_SHA256) -> list[dict]:
    """Bind source members by exact identity; input order must have no semantic effect."""
    by_path_hash = {row["member_path_sha256"]: row for row in observed}
    expected_path_hashes = {contract[2] for contract in contracts}
    if len(by_path_hash) != len(observed) or set(by_path_hash) != expected_path_hashes:
        raise ValueError("unexpected shape member set")
    bound = []
    for member_id, scenario_label, path_sha256, expected_dbf_sha256 in contracts:
        row = by_path_hash[path_sha256]
        if row["shp_sha256"] != shp_sha256:
            raise ValueError(f"unexpected SHP bytes for {member_id}")
        if row["dbf_sha256"] != expected_dbf_sha256:
            raise ValueError(f"unexpected DBF bytes for {member_id}")
        if row["txt_sha256"] != txt_sha256:
            raise ValueError(f"unexpected CRS sidecar bytes for {member_id}")
        bound.append({**row, "member_id": member_id, "scenario_label": scenario_label})
    return bound


def _validated_shape_members(archive: zipfile.ZipFile, filename: str) -> list[dict]:
    """Bind each scenario to an exact member identity, independent of archive order."""
    observed = []
    for shp_name in (name for name in archive.namelist() if name.endswith(".shp")):
        base = shp_name[:-4]
        required = [base + extension for extension in (".shp", ".shx", ".dbf", ".txt")]
        missing = [name for name in required if name not in archive.namelist()]
        if missing:
            raise ValueError(f"missing shape member companions: {missing}")
        observed.append({
            "shp_name": shp_name,
            "member_path_sha256": _digest(_normalized_member_name(archive, shp_name).encode("utf-8")),
            "shp_sha256": _digest(archive.read(base + ".shp")),
            "dbf_sha256": _digest(archive.read(base + ".dbf")),
            "txt_sha256": _digest(archive.read(base + ".txt")),
        })
    if filename == "01_.zip":
        return _bind_hazard_member_inventory(observed, HAZARD_MEMBER_CONTRACTS[filename], shp_sha256=_COMMON_INTENSITY_SHP_SHA256, txt_sha256=_COMMON_INTENSITY_TXT_SHA256)
    return _bind_hazard_member_inventory(observed, HAZARD_MEMBER_CONTRACTS[filename])


def _validate_definition_workbook(payload: bytes) -> dict[str, list[str]]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    validated = {}
    for source in HAZARD_SOURCES.values():
        if source["definition_sheet"] not in workbook.sheetnames:
            raise ValueError(f"missing definition sheet: {source['definition_sheet']}")
        sheet = workbook[source["definition_sheet"]]
        first_column = [row[0] for row in sheet.iter_rows(values_only=True) if row and isinstance(row[0], str)]
        definition_fields = source.get("definition_fields", source["fields"])
        missing = [field for field in definition_fields if field not in first_column]
        if missing:
            raise ValueError(f"definition workbook lacks fields for {source['kind']}: {missing}")
        validated[source["kind"]] = definition_fields
    return validated


def _build_hazards(repo: Path, raw_root: Path) -> dict:
    license_path = repo / "inputs/staging/FUJISAWA-INTENSITY-CRS-V1/license_binding.json"
    license_bytes = license_path.read_bytes().replace(b"\r\n", b"\n")
    license_binding = json.loads(license_bytes)
    if (license_binding.get("resource_id") != HAZARD_SOURCES["01_.zip"]["resource_id"]
        or license_binding.get("retained_raw", {}).get("sha256") != HAZARD_SOURCES["01_.zip"]["sha256"]
        or license_binding.get("status") != "CC-BY-4.0"
        or license_binding.get("scope") != "KANAGAWA_01_INTENSITY_EIGHT_SCENARIOS_ONLY"
        or license_binding.get("license_url") != "https://creativecommons.org/licenses/by/4.0/deed.ja"):
        raise ValueError("intensity license binding identity/scope mismatch")
    coverage = json.loads((repo / "inputs/staging/PHASE4-MULTICITY-DATA-ACQUISITION/city_aoi/aoi_fujisawa_enoshima_v1.geojson").read_text(encoding="utf-8"))["features"][0]
    definition_name, definition_sha, definition_resource = DEFINITION
    definition_bytes = _verified_bytes(raw_root / definition_name, definition_sha)
    validated_definition_fields = _validate_definition_workbook(definition_bytes)
    layers = []
    display_features = []
    sibling_rings, sibling_sidecar_sha256 = _intensity_sibling_rings(raw_root)
    intensity_summary = {"scenario_count": 0, "selected_feature_count": 0, "max_vertex_error_m": 0.0, "max_sibling_vertex_error_m": 0.0}
    for filename, source in HAZARD_SOURCES.items():
        raw_bytes = _verified_bytes(raw_root / filename, source["sha256"])
        archive = zipfile.ZipFile(io.BytesIO(raw_bytes))
        members = _validated_shape_members(archive, filename)
        for layer_index, bound_member in enumerate(members, 1):
            member_id = bound_member["member_id"]
            bound_scenario_label = bound_member["scenario_label"]
            shp_name = bound_member["shp_name"]
            crs_closure = None
            if source["kind"] == "INTENSITY_SCENARIO":
                all_features, member_receipt, projected_rings = _intensity_shape_features(archive, shp_name, bound_member["member_path_sha256"], source["fields"])
                source_crs = "EPSG:6668"
                selected = select_intersecting_source_features(all_features, coverage, source_crs=source_crs, coverage_crs="EPSG:4326")
                to_projected = Transformer.from_crs(source_crs, "EPSG:6677", always_xy=True, allow_ballpark=False).transform
                max_vertex_error_m = 0.0
                max_sibling_error_m = 0.0
                for feature in selected:
                    mesh_code = feature["properties"]["mesh_code"]
                    corners = [to_projected(*vertex) for vertex in feature["geometry"]["coordinates"][0][:4]]
                    max_vertex_error_m = max(max_vertex_error_m, _max_vertex_error_m(corners, projected_rings[mesh_code]))
                    if mesh_code not in sibling_rings:
                        raise ValueError(f"provider sibling layer lacks selected mesh cell: {mesh_code}")
                    max_sibling_error_m = max(max_sibling_error_m, _max_vertex_error_m(sibling_rings[mesh_code], projected_rings[mesh_code]))
                _assert_intensity_cross_check(max_vertex_error_m, max_sibling_error_m, member_id)
                crs_closure = {
                    "decision": "EPSG:6677_JGD2011_ZONE_IX_PROVIDER_DECLARED",
                    "display_geometry_source": "JIS_X_0410_MESH_CODE_EPSG_6668",
                    "basis": [
                        "provider sibling layer 00_ sidecar PROJCS JGD_2011_Japan_Zone_9 + identical geometry",
                        "record LON/LAT = JIS X 0410 quarter-mesh centre (definition workbook: LON=経度, LAT=緯度)",
                        f"EPSG:6677 cross-check max vertex error <= {INTENSITY_CROSS_CHECK_TOLERANCE_M} m",
                    ],
                    "tolerance_m": INTENSITY_CROSS_CHECK_TOLERANCE_M,
                    "centre_tolerance_m": INTENSITY_CENTRE_TOLERANCE_M,
                    "max_centre_error_m": max(_intensity_centre_error_m(f["properties"]["mesh_code"], f["properties"]["longitude_raw"], f["properties"]["latitude_raw"]) for f in all_features),
                    "tolerance_scope": "VERIFICATION_ONLY_NOT_REGISTRY_NOT_EVALUATION",
                    "tolerance_rationale": "T-B METHOD#3: GRS80 centre distance <= 0.06 m for every source record; METHOD#7: EPSG:6677 vertex cross-check <= 0.05 m. Verification only, not measurement accuracy or M7 evaluation.",
                    "license_review": "CC-BY-4.0",
                    "license_note": INTENSITY_LICENSE_NOTE,
                    "license_binding_sha256": _digest(license_bytes),
                    "max_vertex_error_m": max_vertex_error_m,
                    "max_sibling_vertex_error_m": max_sibling_error_m,
                    "sibling_source_sha256": INTENSITY_SIBLING["sha256"],
                    "sibling_sidecar_sha256": sibling_sidecar_sha256,
                    "scope": "APPLIES_ONLY_TO_01_INTENSITY_SCENARIO_MEMBERS",
                    "erroneous_sidecar": {
                        "text_sha256": member_receipt["txt_sha256"],
                        "declared": INTENSITY_ERRONEOUS_SIDECAR_DECLARED,
                        "status": "ERRONEOUS_SIDECAR_RECORDED_NOT_USED",
                    },
                }
                intensity_summary["scenario_count"] += 1
                intensity_summary["selected_feature_count"] += len(selected)
                intensity_summary["max_vertex_error_m"] = max(intensity_summary["max_vertex_error_m"], max_vertex_error_m)
                intensity_summary["max_sibling_vertex_error_m"] = max(intensity_summary["max_sibling_vertex_error_m"], max_sibling_error_m)
            else:
                all_features, member_receipt = _shape_features(archive, shp_name, bound_member["member_path_sha256"], source["kind"], source["fields"])
                source_crs = "EPSG:4612"
                selected = select_intersecting_source_features(all_features, coverage, source_crs=source_crs, coverage_crs="EPSG:4326")
            scenario_suffix = f"_{layer_index:02d}" if len(members) > 1 else ""
            if source["kind"] == "INTENSITY_SCENARIO":
                scenario_id = f"fujisawa_earthquake_intensity_scenario_{layer_index:02d}"
                scenario_label = bound_scenario_label
                layer_kind = "震度分布"
            elif source["kind"] == "LIQUEFACTION_SCENARIO":
                scenario_id = f"fujisawa_liquefaction_distribution_scenario_{layer_index:02d}"
                scenario_label = bound_scenario_label
                layer_kind = "液状化分布"
            elif source["kind"] == "SHAKING_SUSCEPTIBILITY":
                scenario_id = "fujisawa_shaking_susceptibility_r6_01"
                scenario_label = bound_scenario_label
                layer_kind = "ゆれやすさ"
            else:
                scenario_id = "fujisawa_liquefaction_hazard_r6_01"
                scenario_label = bound_scenario_label
                layer_kind = "液状化危険度"
            layer_source_id = source["source_id"] + scenario_suffix.lower()
            selection_payload = json.dumps(selected, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
            layer = {
                "source_id": layer_source_id, "source_revision": source["member_revision"],
                "dataset_publication_date": "2025-03-28",
                "scenario_label": scenario_label, "layer_kind": layer_kind,
                "source_sha256": source["sha256"],
                "source_url": f"{KANAGAWA_DATASET}/resource/{source['resource_id']}",
                "resource_response_sha256": source["resource_response_sha256"],
                # Current metadata binding supersedes the historical pending receipt,
                # only for this resource, never Fujisawa municipal facility records.
                "license_status": "CC-BY-4.0" if crs_closure else "CC-BY",
                **({"license_note": INTENSITY_LICENSE_NOTE} if crs_closure else {}),
                "license_url": license_binding["license_url"] if crs_closure else KANAGAWA_DATASET,
                "attribution": "神奈川県『地震被害想定調査（令和7年3月）』を加工して作成",
                "definition_resource_id": definition_resource,
                "definition_resource_response_sha256": KANAGAWA_DEFINITION_RESPONSE_SHA256,
                "definition_sha256": definition_sha, "definition_fields": validated_definition_fields[source["kind"]],
                "source_member_id": member_id,
                "source_member_receipt": member_receipt,
                "crs_sidecar_sha256": member_receipt["txt_sha256"], "dbf_encoding": "CP932_EXPLICIT",
                "scenario_id": scenario_id,
                "source_crs": source_crs, "coverage": coverage, "coverage_crs": "EPSG:4326",
                "coverage_evidence": {"status": "FULL_SOURCE_SCAN_INTERSECTED_WITH_BOUND_AOI", "selection_sha256": _digest(selection_payload)},
                "class_field": "source_class", "features": selected,
                "limitations": "Raw official mesh attributes (JMA/PGA/SI) and geometric overlap only; HAZARD_EXPOSURE_ONLY; no threshold, closure, damage, debris, passability, accessibility, or safety state is derived; scenarios are not averaged." if crs_closure else "Raw official mesh attributes and geometric overlap only; no threshold, closure, damage, debris, passability, accessibility, or safety state is derived.",
                **({"crs_closure": crs_closure} if crs_closure else {}),
            }
            layers.append(layer)
            to_wgs84 = Transformer.from_crs(source_crs, "EPSG:4326", always_xy=True, allow_ballpark=False).transform
            for feature in selected:
                display_features.append({"type": "Feature", "geometry": mapping(transform(to_wgs84, shape(feature["geometry"]))), "properties": {**feature["properties"], "source_id": layer_source_id, "scenario_id": scenario_id, **({"_ablepath_display_only": True, "_ablepath_exposure_only": True} if crs_closure else {})}})
    return {
        "schema_version": "1.0.0", "dataset_url": KANAGAWA_DATASET,
        "dataset_response_sha256": KANAGAWA_DATASET_RESPONSE_SHA256,
        "license_status": "CC-BY", "license_url": KANAGAWA_DATASET,
        "attribution": "神奈川県『地震被害想定調査（令和7年3月）』を加工して作成",
        "raw_root_id": "EXTERNAL_OFFICIAL_DATA_ROOT",
        "intensity_distribution": {
            "status": "CONNECTED",
            "source_sha256": HAZARD_SOURCES["01_.zip"]["sha256"],
            "scenario_count": intensity_summary["scenario_count"],
            "selected_feature_count": intensity_summary["selected_feature_count"],
            "crs_closure": {
                "decision": "EPSG:6677_JGD2011_ZONE_IX_PROVIDER_DECLARED",
                "display_geometry_source": "JIS_X_0410_MESH_CODE_EPSG_6668",
                "tolerance_m": INTENSITY_CROSS_CHECK_TOLERANCE_M,
                "centre_tolerance_m": INTENSITY_CENTRE_TOLERANCE_M,
                "max_centre_error_m": max(layer["crs_closure"]["max_centre_error_m"] for layer in layers if "crs_closure" in layer),
                "tolerance_scope": "VERIFICATION_ONLY_NOT_REGISTRY_NOT_EVALUATION",
                "max_vertex_error_m": intensity_summary["max_vertex_error_m"],
                "max_sibling_vertex_error_m": intensity_summary["max_sibling_vertex_error_m"],
                "sibling_source_sha256": INTENSITY_SIBLING["sha256"],
                "scope": "APPLIES_ONLY_TO_01_INTENSITY_SCENARIO_MEMBERS",
                "erroneous_sidecar_status": "ERRONEOUS_SIDECAR_RECORDED_NOT_USED",
            },
            "exposure_only": True, "scenario_averaged": False,
            "license_review": "CC-BY-4.0", "license_note": INTENSITY_LICENSE_NOTE,
            "license_binding_sha256": _digest(license_bytes),
            "reason": "The provider's own sibling-layer PROJCS declaration and the JIS X 0410 mesh-code centres close the CRS question; display geometry is rebuilt from KEY_CODE and the projected SHP is used only as a cross-check. Exposure display only.",
        },
        "layers": layers, "display_features": display_features,
        "closure_derived": False, "damage_or_debris_inferred": False,
    }


REUSED_DEM_REQUIRED_KEYS = ("cities", "outer_sha256", "no_interpolation", "no_smoothing", "no_step_inference", "no_cross_slope_inference")


def _reused_dem_section(source: Path) -> tuple[dict, dict]:
    """Copy the `dem` section verbatim from an existing evidence file, fail-closed.

    The DEM raw package is explicitly not re-acquired (PR14 freeze receipt).
    Nothing is synthesized, and an arbitrary or tampered file is rejected rather
    than silently promoted into the public evidence payload.
    """
    source_bytes = source.read_bytes()
    payload = json.loads(source_bytes.decode("utf-8"))
    if not isinstance(payload, dict) or not isinstance(payload.get("dem"), dict):
        raise ValueError(f"reuse source has no `dem` section: {source.name}")
    dem = payload["dem"]
    missing = [key for key in REUSED_DEM_REQUIRED_KEYS if key not in dem]
    if missing:
        raise ValueError(f"reuse source `dem` section lacks required keys: {missing}")
    if dem["outer_sha256"] != OUTER_DEM[1]:
        raise ValueError("reuse source `dem` section is bound to a different DEM raw package")
    if not isinstance(dem["cities"], dict) or set(dem["cities"]) != set(DEM_PRODUCTS):
        raise ValueError("reuse source `dem` section does not cover exactly the bound cities")
    if any(dem[key] is not True for key in ("no_interpolation", "no_smoothing", "no_step_inference", "no_cross_slope_inference")):
        raise ValueError("reuse source `dem` section does not retain the no-inference contract")
    section_sha = _digest(json.dumps(dem, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))
    if section_sha != FROZEN_DEM_SECTION_SHA256:
        raise ValueError("reuse source differs from the frozen DEM section")
    return dem, {
        "mode": "REUSED_VERBATIM_FROM_COMMITTED_EVIDENCE",
        "source_sha256": _digest(source_bytes),
        "section_sha256": section_sha,
        "section_hash_format": "UTF8_JSON_SORT_KEYS_COMPACT",
        "reason": "DEM re-acquisition explicitly not repeated (PR14 freeze receipt); frozen DEM section verified without re-reading the DEM raw package",
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", type=Path, required=True)
    parser.add_argument("--raw-root", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--reuse-dem-from", type=Path, help="Copy the `dem` section verbatim from an existing official_evidence.json instead of re-reading the DEM raw package.")
    args = parser.parse_args()
    output = args.output or args.repo_root / "inputs/staging/PUBLIC-GIT-DEM-FUJISAWA-V1/official_evidence.json"
    if args.reuse_dem_from:
        dem, dem_rebuild = _reused_dem_section(args.reuse_dem_from)
    else:
        dem = _build_dem(args.repo_root, args.raw_root)
        dem_rebuild = {"mode": "REBUILT_FROM_VERIFIED_RAW_PACKAGE", "source_sha256": OUTER_DEM[1], "reason": "The verified DEM raw package was present in this run."}
    payload = {"dem": dem, "dem_rebuild": dem_rebuild, "fujisawa_hazards": _build_hazards(args.repo_root, args.raw_root)}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps({"output": str(output), "sha256": _digest(output.read_bytes())}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
